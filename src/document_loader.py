"""Document Loader: extract, clean and de-duplicate text from PDF, DOCX, TXT and Markdown."""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Union

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}

# DOCX files have no real pages; we group paragraphs into "sections" of roughly this many words
DOCX_WORDS_PER_SECTION = 450


@dataclass
class Page:
    """One logical unit of a document (a PDF page, a DOCX section or a text block)."""

    text: str
    page_number: int  # 1-based
    heading: str = ""


@dataclass
class LoadedDocument:
    file_name: str
    file_type: str
    pages: List[Page]
    content_hash: str
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    @property
    def num_pages(self) -> int:
        return len(self.pages)

    @property
    def word_count(self) -> int:
        return sum(len(p.text.split()) for p in self.pages)


class UnsupportedFileError(ValueError):
    pass


# --------------------------------------------------------------------------- cleaning
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_MULTI_SPACE = re.compile(r"[ \t ]+")
_MULTI_NEWLINE = re.compile(r"\n{3,}")
_PAGE_NUMBER_LINE = re.compile(r"^\s*(page\s*)?\d{1,4}\s*(of\s*\d{1,4})?\s*$", re.IGNORECASE)


def clean_text(text: str) -> str:
    """Normalise unicode, fix hyphenated line breaks, collapse whitespace, drop bare page-number lines."""
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    lines = []
    for line in text.split("\n"):
        line = _MULTI_SPACE.sub(" ", line).strip()
        if _PAGE_NUMBER_LINE.match(line):
            continue
        lines.append(line)
    text = "\n".join(lines)
    text = _MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


def remove_duplicate_paragraphs(pages: List[Page], min_len: int = 40) -> List[Page]:
    """Remove repeated paragraphs (running headers/footers, copy-pasted blocks) across the document."""
    seen: set[str] = set()
    out: List[Page] = []
    for page in pages:
        kept = []
        for para in page.text.split("\n\n"):
            key = re.sub(r"\W+", " ", para.lower()).strip()
            if len(key) >= min_len:
                if key in seen:
                    continue
                seen.add(key)
            kept.append(para)
        new_text = "\n\n".join(kept).strip()
        if new_text:
            out.append(Page(text=new_text, page_number=page.page_number, heading=page.heading))
    return out


def remove_repeated_lines(pages: List[Page], min_share: float = 0.4, edge: int = 3) -> List[Page]:
    """Drop running headers/footers: short lines at the top or bottom of a page that repeat on a large share
    of pages (e.g. "Registration no.: Branch & Section: Faculty Name:" on every page of a lab manual).
    Lines in the middle of a page (such as repeated code like `import java.util.*;`) are never touched."""
    if len(pages) < 3:
        return pages

    def key(line: str) -> str:
        return re.sub(r"[\d\W_]+", " ", line.lower()).strip()

    def edges(lines: List[str]) -> set:
        idx = [i for i, l in enumerate(lines) if l.strip()]
        return set(idx[:edge] + idx[-edge:])

    counts: dict[str, int] = {}
    for p in pages:
        lines = p.text.split("\n")
        for k in {key(lines[i]) for i in edges(lines) if len(lines[i].strip()) <= 140}:
            if k:
                counts[k] = counts.get(k, 0) + 1
    limit = max(3, int(len(pages) * min_share))
    repeated = {k for k, c in counts.items() if c >= limit}
    if not repeated:
        return pages
    out = []
    for p in pages:
        lines = p.text.split("\n")
        drop = {i for i in edges(lines) if len(lines[i].strip()) <= 140 and key(lines[i]) in repeated}
        text = "\n".join(l for i, l in enumerate(lines) if i not in drop).strip()
        if text:
            out.append(Page(text=text, page_number=p.page_number, heading=p.heading))
    return out


def compute_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --------------------------------------------------------------------------- loaders
def _load_pdf(data: bytes) -> List[Page]:
    pages: List[Page] = []
    try:
        import fitz  # PyMuPDF – fastest and most accurate

        with fitz.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, start=1):
                pages.append(Page(text=page.get_text("text"), page_number=i))
        return pages
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # fallback
    except ImportError as e:
        raise ImportError(
            "No PDF reader is installed in the Python that is running the app. Close the app, then in the project "
            "folder run:  .venv\\Scripts\\activate  then  python -m pip install -r requirements-lite.txt  and start it "
            "again with:  python -m streamlit run app.py"
        ) from e

    reader = PdfReader(io.BytesIO(data))
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        pages.append(Page(text=txt, page_number=i))
    return pages


def _load_docx(data: bytes) -> List[Page]:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    blocks: List[tuple[str, str]] = []  # (text, style)
    for para in document.paragraphs:
        if para.text.strip():
            blocks.append((para.text.strip(), (para.style.name or "") if para.style else ""))
    # include table content
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append((" | ".join(dict.fromkeys(cells)), "Table"))

    pages: List[Page] = []
    current: List[str] = []
    words = 0
    heading = ""
    current_heading = ""
    for text, style in blocks:
        is_heading = style.lower().startswith("heading") or style.lower() == "title"
        # start a new section at a heading once the current one has some content
        if (is_heading and words > DOCX_WORDS_PER_SECTION * 0.5) or words >= DOCX_WORDS_PER_SECTION:
            pages.append(Page(text="\n\n".join(current), page_number=len(pages) + 1, heading=current_heading))
            current, words = [], 0
            current_heading = heading
        if is_heading:
            heading = text
            if not current:
                current_heading = text
        current.append(text)
        words += len(text.split())
    if current:
        pages.append(Page(text="\n\n".join(current), page_number=len(pages) + 1, heading=current_heading))
    return pages


def _load_text(data: bytes, markdown: bool = False) -> List[Page]:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = data.decode("utf-8", errors="ignore")

    if markdown:
        # Split on top-level headings so citations can point to a section
        sections = re.split(r"(?m)^(?=#{1,2}\s)", text)
    else:
        sections = [text]

    pages: List[Page] = []
    for sec in sections:
        if not sec.strip():
            continue
        heading = ""
        m = re.match(r"#{1,6}\s+(.*)", sec.strip())
        if m:
            heading = m.group(1).strip()
        # very long plain-text files are split into ~pages of 3000 chars on paragraph boundaries
        paras: List[str] = []
        for para in re.split(r"\n\s*\n", sec):
            if sum(len(x) for x in paras) + len(para) > 3000 and paras:
                carry = []
                # don't leave a heading stranded at the bottom of a page
                if len(paras) > 1 and len(paras[-1]) < 80 and not re.search(r"[.!?]$", paras[-1].strip()):
                    carry = [paras.pop()]
                pages.append(Page(text="\n\n".join(paras), page_number=len(pages) + 1, heading=heading))
                paras = carry
            paras.append(para)
        buf = "\n\n".join(paras)
        if buf.strip():
            pages.append(Page(text=buf, page_number=len(pages) + 1, heading=heading))
    return pages


def load_document(source: Union[str, Path, bytes], file_name: Optional[str] = None) -> LoadedDocument:
    """Load a document from a path or raw bytes.

    Returns a LoadedDocument with cleaned, de-duplicated pages and a content hash
    (used to detect the same file being uploaded twice).
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        data = path.read_bytes()
        file_name = file_name or path.name
    else:
        data = source
        if not file_name:
            raise ValueError("file_name is required when loading from bytes")

    ext = Path(file_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")

    if ext == ".pdf":
        raw_pages = _load_pdf(data)
    elif ext == ".docx":
        raw_pages = _load_docx(data)
    else:
        raw_pages = _load_text(data, markdown=ext in {".md", ".markdown"})

    pages = [Page(text=clean_text(p.text), page_number=p.page_number, heading=p.heading) for p in raw_pages]
    pages = [p for p in pages if p.text]
    pages = remove_repeated_lines(pages)
    pages = remove_duplicate_paragraphs(pages)

    return LoadedDocument(
        file_name=file_name,
        file_type=ext.lstrip("."),
        pages=pages,
        content_hash=compute_hash(data),
        metadata={"size_bytes": len(data)},
    )
